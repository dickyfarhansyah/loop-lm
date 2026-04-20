from __future__ import annotations
import json
import logging
import os
import io
import threading
from datetime import datetime
from pathlib import Path
from nanoid import generate
from sqlalchemy.orm import Session
from src.config.database import SessionLocal
from src.config.env import env
from src.db.models import File, ChatFile
from src.utils.errors import NotFoundError, BadRequestError
from src.services.chroma_service import chroma_service

logger = logging.getLogger(__name__)

SUPPORTED_MIMETYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/plain": "txt",
    "text/csv": "csv",
}


def _serialize(file: File) -> dict:
    return {
        "id": file.id,
        "userId": file.user_id,
        "filename": file.filename,
        "path": file.path,
        "meta": json.loads(file.meta) if isinstance(file.meta, str) and file.meta else None,
        "createdAt": file.created_at.isoformat() if file.created_at else None,
        "updatedAt": file.updated_at.isoformat() if file.updated_at else None,
    }


def _extract_text(data: bytes, mimetype: str) -> str:
    """
    Extract plain text from various file formats.
    Uses PyMuPDF (fitz) for PDFs — 10-20x faster than pypdf.
    """
    fmt = SUPPORTED_MIMETYPES.get(mimetype, "")

    if fmt == "pdf":
        # PyMuPDF: C-level PDF renderer, much faster than pypdf
        import fitz  # pymupdf
        doc = fitz.open(stream=data, filetype="pdf")
        pages = []
        for page in doc:
            # get_text("text") is fast; "blocks" option gives layout-aware text
            pages.append(page.get_text("text"))
        doc.close()
        return "\n".join(pages)

    elif fmt == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)

    elif fmt == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                lines.append("\t".join(str(c) if c is not None else "" for c in row))
        wb.close()
        return "\n".join(lines)

    elif fmt in ("txt", "csv"):
        return data.decode("utf-8", errors="replace")

    return ""


def _split_text(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
    """
    Split text with LangChain's RecursiveCharacterTextSplitter.
    Uses paragraph → sentence → word boundary hierarchy (same as article recommendation).
    """
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", "! ", "? ", " ", ""],
        length_function=len,
    )
    return splitter.split_text(text)


def _do_embed_in_background(file_id: str, user_id: str, filename: str, fmt: str, text: str) -> None:
    """
    Background thread: chunk the already-extracted text and store embeddings in ChromaDB.
    Text extraction already happened synchronously before this — only the slow embedding is here.
    """
    db = SessionLocal()
    try:
        logger.info(f"[RAG] Background embedding start: {file_id} ({filename}), {len(text)} chars")

        if not text.strip():
            _update_status(db, file_id, "error", error="No text extracted")
            return

        # Chunk with RecursiveCharacterTextSplitter
        chunks = _split_text(text, env.CHUNK_SIZE, env.CHUNK_OVERLAP)
        if not chunks:
            _update_status(db, file_id, "error", error="No chunks generated")
            return

        logger.info(f"[RAG] {len(chunks)} chunks, embedding now...")

        # Embed + store in ChromaDB
        chroma_service.store_chunks(
            file_id=file_id,
            user_id=user_id,
            filename=filename,
            fmt=fmt,
            chunks=[{"text": c, "index": i} for i, c in enumerate(chunks)],
        )

        logger.info(f"[RAG] Stored {len(chunks)} chunks for {file_id}")
        _update_status(db, file_id, "processed", text_length=len(text), chunk_count=len(chunks))

    except Exception as e:
        logger.error(f"[RAG] Embedding failed for {file_id}: {e}", exc_info=True)
        _update_status(db, file_id, "error", error=str(e))
    finally:
        db.close()


def _do_parse_in_background(file_id: str, user_id: str, filename: str, fmt: str, data: bytes, mimetype: str) -> None:
    """Legacy — kept for compatibility. Calls _do_embed_in_background after extracting text."""
    text = _extract_text(data, mimetype)
    _do_embed_in_background(file_id, user_id, filename, fmt, text)


def _update_status(db: Session, file_id: str, status: str, **extra) -> None:
    """Update file meta.status in DB."""
    f = db.query(File).filter(File.id == file_id).first()
    if not f:
        return
    meta = json.loads(f.meta) if isinstance(f.meta, str) and f.meta else {}
    meta["status"] = status
    meta.update(extra)
    f.meta = json.dumps(meta)
    f.updated_at = datetime.utcnow()
    db.commit()


class FileService:
    def _save_file(self, data: bytes, filename: str, user_id: str) -> str:
        """Save bytes to disk and return the relative path."""
        user_dir = Path(env.UPLOAD_DIR) / user_id
        user_dir.mkdir(parents=True, exist_ok=True)
        file_id = generate()
        safe_name = Path(filename).name
        dest = user_dir / f"{file_id}_{safe_name}"
        dest.write_bytes(data)
        return str(dest)

    def upload_file(self, db: Session, file_info: dict, user_id: str) -> dict:
        now = datetime.utcnow()
        file_id = generate()
        path = self._save_file(file_info["data"], file_info["filename"], user_id)
        meta = json.dumps({
            "mimetype": file_info["mimetype"],
            "size": len(file_info["data"]),
        })
        f = File(id=file_id, user_id=user_id, filename=file_info["filename"],
                 path=path, meta=meta, created_at=now, updated_at=now)
        db.add(f)
        db.commit()
        db.refresh(f)
        return _serialize(f)

    def get_files(self, db: Session, user_id: str) -> list:
        files = db.query(File).filter(File.user_id == user_id).all()
        return [_serialize(f) for f in files]

    def get_file_by_id(self, db: Session, file_id: str, user_id: str) -> dict:
        f = db.query(File).filter(File.id == file_id, File.user_id == user_id).first()
        if not f:
            raise NotFoundError("File not found")
        return _serialize(f)

    def get_file_status(self, db: Session, file_id: str, user_id: str) -> dict:
        """Poll processing status for a file."""
        f = db.query(File).filter(File.id == file_id, File.user_id == user_id).first()
        if not f:
            raise NotFoundError("File not found")
        meta = json.loads(f.meta) if isinstance(f.meta, str) and f.meta else {}
        return {
            "id": f.id,
            "filename": f.filename,
            "status": meta.get("status", "unknown"),
            "textLength": meta.get("text_length"),
            "chunkCount": meta.get("chunk_count"),
            "error": meta.get("error"),
        }

    def download_file(self, db: Session, file_id: str, user_id: str) -> tuple:
        f = db.query(File).filter(File.id == file_id, File.user_id == user_id).first()
        if not f:
            raise NotFoundError("File not found")
        meta = json.loads(f.meta) if isinstance(f.meta, str) and f.meta else {}
        data = Path(f.path).read_bytes()
        return data, f.filename, meta

    def download_file_public(self, db: Session, file_id: str) -> tuple:
        """Serve a file without user check — only used for public image assets like logos."""
        f = db.query(File).filter(File.id == file_id).first()
        if not f:
            raise NotFoundError("File not found")
        meta = json.loads(f.meta) if isinstance(f.meta, str) and f.meta else {}
        data = Path(f.path).read_bytes()
        return data, f.filename, meta

    def delete_file(self, db: Session, file_id: str, user_id: str) -> None:
        f = db.query(File).filter(File.id == file_id, File.user_id == user_id).first()
        if not f:
            raise NotFoundError("File not found")
        try:
            os.remove(f.path)
        except OSError:
            pass
        chroma_service.delete_file_chunks(file_id)
        db.delete(f)
        db.commit()

    def attach_file_to_chat(self, db: Session, chat_id: str, file_id: str, user_id: str, message_id: str | None = None) -> dict:
        now = datetime.utcnow()
        existing = db.query(ChatFile).filter(
            ChatFile.chat_id == chat_id, ChatFile.file_id == file_id
        ).first()
        if existing:
            return {"chatId": chat_id, "fileId": file_id}
        cf = ChatFile(
            id=generate(), user_id=user_id, chat_id=chat_id,
            file_id=file_id, message_id=message_id,
            created_at=now, updated_at=now,
        )
        db.add(cf)
        db.commit()
        return {"chatId": chat_id, "fileId": file_id, "messageId": message_id}

    def is_supported_document(self, mimetype: str) -> bool:
        return mimetype in SUPPORTED_MIMETYPES

    def upload_and_parse_document(self, db: Session, file_info: dict, user_id: str) -> dict:
        """
        Upload file, extract text synchronously (fast), then embed in background.

        Frontend expects:
          { "file": {...}, "parsed": { "text": "...", "metadata": {...} } }

        Pipeline:
          1. PyMuPDF   → fast text extraction (sync, ~ms)
          2. Return text immediately to frontend
          3. RecursiveCharacterTextSplitter + ChromaDB → background thread
        """
        mimetype = file_info["mimetype"]
        if not self.is_supported_document(mimetype):
            raise BadRequestError(f"Unsupported format: {mimetype}")

        fmt = SUPPORTED_MIMETYPES[mimetype]

        # Step 1: Extract text synchronously (PyMuPDF is fast, ~ms for most PDFs)
        text = _extract_text(file_info["data"], mimetype)

        # Step 2: Save file to disk + DB with status=processing
        now = datetime.utcnow()
        file_id = generate()
        path = self._save_file(file_info["data"], file_info["filename"], user_id)
        word_count = len(text.split()) if text else 0
        char_count = len(text) if text else 0
        meta = json.dumps({
            "mimetype": mimetype,
            "size": len(file_info["data"]),
            "status": "processing",
            "wordCount": word_count,
            "charCount": char_count,
        })
        f = File(id=file_id, user_id=user_id, filename=file_info["filename"],
                 path=path, meta=meta, created_at=now, updated_at=now)
        db.add(f)
        db.commit()

        # Step 3: Chunk + embed in background (slow part — sentence-transformers)
        thread = threading.Thread(
            target=_do_embed_in_background,
            args=(file_id, user_id, file_info["filename"], fmt, text),
            daemon=True,
        )
        thread.start()

        # Step 4: Return text immediately to frontend — matches ParseDocumentResponse shape
        return {
            "file": {
                "id": file_id,
                "filename": file_info["filename"],
                "path": path,
                "meta": {
                    "mimetype": mimetype,
                    "size": len(file_info["data"]),
                },
            },
            "parsed": {
                "text": text,
                "metadata": {
                    "format": fmt,
                    "wordCount": word_count,
                    "charCount": char_count,
                },
            },
        }


file_service = FileService()

